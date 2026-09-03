"""Render Presentation Agent JSON output as a downloadable PowerPoint file."""

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt


def _slides(payload: dict[str, Any]) -> list[dict[str, Any]]:
    body = payload.get("body", payload)
    slides = body.get("slides", []) if isinstance(body, dict) else []
    return slides if isinstance(slides, list) else []


def build_presentation(payload: dict[str, Any], title: str = "Executive Meeting Summary") -> bytes:
    """Build a simple executive presentation from Presentation Agent JSON."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    slides = _slides(payload)
    for slide_data in slides:
        slide = presentation.slides.add_slide(blank_layout)
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.8), Inches(0.7))
        title_frame = title_box.text_frame
        title_frame.text = str(slide_data.get("title", title))
        title_frame.paragraphs[0].font.size = Pt(28)
        title_frame.paragraphs[0].font.bold = True

        bullets = slide_data.get("bullets", [])
        content_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.2), Inches(4.8))
        frame = content_box.text_frame
        frame.word_wrap = True
        for index, bullet in enumerate(bullets if isinstance(bullets, list) else []):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.level = 0
            paragraph.font.size = Pt(20)

        notes = slide_data.get("speaker_notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)

    if not slides:
        slide = presentation.slides.add_slide(blank_layout)
        text_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.0))
        text_box.text_frame.text = title

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
